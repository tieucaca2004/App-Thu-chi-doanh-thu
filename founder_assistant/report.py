"""Founder Daily Report: short text for Zalo + XLSX file (SUMMARY, PURCHASE, SALES, PRICE_CHANGE, ALERTS, SOURCES)."""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analytics import CATEGORY_LABEL, DailySummary, daily_summary
from .db import DB
from .money import fmt_k, fmt_pct, fmt_signed_vnd, fmt_vnd
from .pricing import NOT_ENOUGH_DATA, average_price, trend_statement

LINE = "══════════════════════════"


def render_text(db: DB, s: DailySummary) -> str:
    d = date.fromisoformat(s.day)
    out = [LINE, "     FOUNDER DAILY REPORT", f"        {d:%d/%m/%Y}", LINE, ""]

    out += ["💰 DOANH THU", fmt_vnd(s.revenue) if s.revenue is not None else "UNKNOWN — chưa có dữ liệu"]
    if s.bill_count is not None:
        out.append(f"Số bill: {s.bill_count}")
    if s.revenue_note:
        out.append(f"ℹ️ {s.revenue_note}")
    out.append("")

    out.append("💸 CHI")
    out.append(fmt_vnd(s.expense) if s.expense is not None else "UNKNOWN — chưa có dữ liệu")
    for cat in CATEGORY_LABEL:
        if s.expense_by_category.get(cat):
            out.append(f"  {CATEGORY_LABEL[cat]}: {fmt_vnd(s.expense_by_category[cat])}")
    out.append("")

    out.append("📈 THU - CHI")
    out.append(fmt_signed_vnd(s.net) if s.net is not None else "UNKNOWN (thiếu thu hoặc chi)")
    out.append("(Chênh lệch thu – chi, chưa phải lợi nhuận ròng.)")
    out.append("")

    out.append("🛒 GIÁ HÀNG")
    ups = [c for c in s.price_changes if c.diff > 0]
    downs = [c for c in s.price_changes if c.diff < 0]
    if not s.price_changes:
        out.append("Không có thay đổi giá so với lần mua trước." if s.first_prices or s.purchase_rows
                   else "Không có dữ liệu giá hôm nay.")
    for c in sorted(ups, key=lambda c: -c.pct):
        out += [f"⚠️ {c.product_name}", f"{fmt_k(c.old_price)} → {fmt_k(c.new_price)}/{c.base_unit}", fmt_pct(c.pct)]
        trend = trend_statement(db, c.product_id, c.base_unit, s.day)
        if trend != NOT_ENOUGH_DATA:
            out.append(trend)
    if downs:
        out.append("")
    for c in sorted(downs, key=lambda c: c.pct):
        out += [f"↓ {c.product_name}", f"{fmt_k(c.old_price)} → {fmt_k(c.new_price)}/{c.base_unit}", fmt_pct(c.pct)]
    if s.first_prices:
        out.append(f"({len(s.first_prices)} mặt hàng mới có giá lần đầu — chưa có dữ liệu để so sánh.)")
    out.append("")

    out.append("🔥 BÁN CHẠY")
    if s.top_items:
        for name, qty, _ in s.top_items:
            out.append(f"{name}: {qty:g} phần")
    else:
        out.append("Chưa có dữ liệu món bán.")
    out.append("")

    out.append("🚨 CẦN CHÚ Ý")
    attention = [a["message"] for a in s.alerts if a["severity"] in ("critical", "warning")]
    attention += [f"{p['ref']} chờ xác nhận: {p['status_reason']}" for p in s.pending]
    attention += s.missing
    if attention:
        out += [f"- {m}" for m in _dedupe(attention)[:12]]
    else:
        out.append("- Không có.")
    out += ["", LINE]
    return "\n".join(out)


def _dedupe(items: list[str]) -> list[str]:
    seen, out = set(), []
    for i in items:
        if i not in seen:
            seen.add(i)
            out.append(i)
    return out


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = '#,##0" đ"'


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], money_cols: tuple[int, ...] = (), first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for col in money_cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for c in row:
                if isinstance(c.value, (int, float)):
                    c.number_format = MONEY
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"
    return ws


def _u(v):
    """Blank numeric source values are written as UNKNOWN, never as 0."""
    return "UNKNOWN" if v is None else v


def build_xlsx(db: DB, s: DailySummary, path: Path) -> Path:
    wb = Workbook()
    summary = [
        ["Ngày", s.day, ""],
        ["Doanh thu", _u(s.revenue), ", ".join(s.revenue_refs)],
        ["Tổng chi", _u(s.expense), ", ".join(s.expense_refs)],
    ]
    for cat, v in s.expense_by_category.items():
        summary.append([f"  Chi — {CATEGORY_LABEL[cat]}", v, ""])
    summary += [
        ["Thu - chi (chưa phải lợi nhuận ròng)", _u(s.net), "Doanh thu − Tổng chi"],
        ["Số bill", _u(s.bill_count), ""],
        ["Số giao dịch", s.transaction_count, "Số chứng từ đã xác nhận"],
        ["Tiền mặt", _u(s.payments.get("cash")), ""],
        ["Chuyển khoản", _u(s.payments.get("bank_transfer")), ""],
        ["Ví điện tử", _u(s.payments.get("e_wallet")), ""],
        ["Phí nền tảng", _u(s.payments.get("platform_fee")), ""],
    ]
    if s.revenue_note:
        summary.append(["Ghi chú doanh thu", s.revenue_note, ""])
    rec = s.extra.get("reconciliation")
    if rec:
        summary += [
            ["Nguồn doanh thu chính", rec["primary"] or "CHƯA XÁC ĐỊNH", "PRIMARY_REVENUE_SOURCE"],
            ["Đối chiếu — POS / chốt ca", _u(rec["pos_total"]), ", ".join(rec["pos_refs"])],
            ["Đối chiếu — tổng bill lẻ", _u(rec["bill_total"]), ", ".join(rec["bill_refs"])],
            ["Đối chiếu — chênh lệch", _u(rec["difference"]), "KHỚP" if rec["matched"] else "KHÔNG KHỚP — cần xác nhận"],
        ]
    _sheet(wb, "SUMMARY", ["Chỉ số", "Giá trị", "Nguồn"], summary, money_cols=(2,), first=True)

    _sheet(wb, "PURCHASE",
           ["Ngày", "Tên hàng", "Tên trên chứng từ", "Nhóm", "Số lượng", "Đơn vị", "Đơn giá", "Đơn giá (nguồn)",
            "Thành tiền", "Thành tiền (nguồn)", "Nhà cung cấp", "Trạng thái", "Mã chứng từ", "Tin nhắn", "Trích nguồn"],
           [[r["date"], r["name"], r["raw_name"], r["category"], _u(r["quantity"]), r["unit"] or "UNKNOWN",
             _u(r["unit_price"]), _src(r["unit_price_source"]), _u(r["amount"]), _src(r["amount_source"]),
             r["supplier"] or "", r["status"] if r["status"] == "ok" else f"{r['status']}: {r['status_reason']}",
             r["ref"], f"msg#{r['message_id']}", r["evidence"]] for r in s.purchase_rows],
           money_cols=(7, 9))

    _sheet(wb, "SALES",
           ["Ngày", "Món", "Số lượng", "Đơn giá", "Thành tiền", "Tính vào doanh thu", "Mã chứng từ", "Tin nhắn", "Trích nguồn"],
           [[r["date"], r["name"], _u(r["quantity"]), _u(r["unit_price"]), _u(r["amount"]),
             "Có" if r["counted"] else "Không — chỉ đối chiếu (không phải nguồn doanh thu chính)", r["ref"], f"msg#{r['message_id']}", r["evidence"]]
            for r in s.sale_rows],
           money_cols=(4, 5))

    pc_rows = []
    for c in s.price_changes:
        old_src = _price_source(db, c.old_ref)
        new_src = _price_source(db, c.new_ref)
        pc_rows.append([c.product_name, c.base_unit, c.old_price, c.new_price, c.diff, round(c.pct, 2),
                        "Tăng" if c.diff > 0 else "Giảm", c.old_date, old_src, new_src,
                        trend_statement(db, c.product_id, c.base_unit, s.day),
                        *_averages(db, c.product_id, c.base_unit, s.day)])
    for f in s.first_prices:
        pc_rows.append([f["name"], f["unit"], "", f["price"], "", "", "Lần đầu", "", "", "", NOT_ENOUGH_DATA, "", ""])
    _sheet(wb, "PRICE_CHANGE",
           ["Tên hàng", "Đơn vị", "Giá cũ", "Giá mới", "Chênh lệch", "%", "Trạng thái", "Ngày giá cũ",
            "Nguồn giá cũ", "Nguồn giá mới", "Nhận định", "TB 7 ngày", "TB 30 ngày"],
           pc_rows, money_cols=(3, 4, 5, 12, 13))

    sev = {"critical": "NGHIÊM TRỌNG", "warning": "CẢNH BÁO", "info": "THÔNG TIN"}
    alert_rows = [[sev.get(a["severity"], a["severity"]), a["message"], a["code"],
                   f"msg#{a['message_id']}" if a["message_id"] else ""] for a in s.alerts]
    alert_rows += [["CHỜ XÁC NHẬN", f"{p['ref']} ({p['date']}): {p['status_reason']}", "PENDING", f"msg#{p['message_id']}"]
                   for p in s.pending]
    alert_rows += [["THIẾU DỮ LIỆU", m, "MISSING", ""] for m in s.missing]
    _sheet(wb, "ALERTS", ["Mức độ", "Nội dung", "Mã", "Nguồn dữ liệu"], alert_rows)

    _sheet(wb, "SOURCES",
           ["Tin nhắn", "Loại", "Thời gian nhận", "Nội dung gốc / transcript", "Tệp gốc", "SHA-256", "AI extraction", "Model"],
           _source_rows(db, s))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _averages(db: DB, product_id: int, base_unit: str, day: str) -> list:
    out = []
    for days in (7, 30):
        avg, n = average_price(db, product_id, base_unit, day, days)
        out.append(round(avg) if avg is not None and n >= 2 else "Không đủ dữ liệu")
    return out


def _src(v):
    return {"source": "ghi trên nguồn", "computed": "hệ thống tính"}.get(v, "UNKNOWN")


def _price_source(db: DB, ph_id: int) -> str:
    r = db.one("""SELECT ph.source, ph.purchase_item_id, pi.purchase_id, ae.message_id FROM price_history ph
                  LEFT JOIN purchase_items pi ON pi.id = ph.purchase_item_id
                  JOIN ai_extractions ae ON ae.id = ph.extraction_id WHERE ph.id = ?""", (ph_id,))
    if r is None:
        return ""
    if r["purchase_id"]:
        return f"M{r['purchase_id']} / msg#{r['message_id']}"
    return f"bảng giá / msg#{r['message_id']}"


def _source_rows(db: DB, s: DailySummary) -> list[list]:
    mids = ({r["message_id"] for r in s.purchase_rows + s.sale_rows}
            | {a["message_id"] for a in s.alerts if a["message_id"]}
            | {p["message_id"] for p in s.pending})
    # every record behind a headline number (a POS closing has no line items but IS the revenue source)
    for ref in s.revenue_refs + s.expense_refs + (s.extra.get("reconciliation") or {}).get("pos_refs", []) \
            + (s.extra.get("reconciliation") or {}).get("bill_refs", []):
        row = db.one(f"SELECT message_id FROM {'purchases' if ref[0] == 'M' else 'sales'} WHERE id = ?", (int(ref[1:]),))
        if row:
            mids.add(row["message_id"])
    for c in s.price_changes:  # both the old and the new price
        for ph in (c.old_ref, c.new_ref):
            row = db.one("SELECT ae.message_id FROM price_history ph JOIN ai_extractions ae ON ae.id = ph.extraction_id "
                         "WHERE ph.id = ?", (ph,))
            if row:
                mids.add(row["message_id"])
    mids = sorted(mids)
    rows = []
    for mid in mids:
        m = db.one("SELECT * FROM messages WHERE id = ?", (mid,))
        if m is None:
            continue
        media = db.one("SELECT * FROM media WHERE message_id = ?", (mid,))
        ex = db.one("SELECT id, model FROM ai_extractions WHERE message_id = ? ORDER BY id DESC", (mid,))
        rows.append([f"msg#{mid}", m["kind"], m["received_at"], m["transcript"] or m["text"] or "",
                     media["file_path"] if media else "", media["sha256"] if media else "",
                     f"extraction#{ex['id']}" if ex else "", ex["model"] if ex else ""])
    return rows


def generate_daily_report(db: DB, day: str, reports_dir: Path, now: datetime) -> tuple[str, Path, DailySummary]:
    s = daily_summary(db, day)
    text = render_text(db, s)
    path = build_xlsx(db, s, Path(reports_dir) / f"Founder-Daily-Report-{day}.xlsx")
    data = {
        "revenue": s.revenue, "expense": s.expense, "net": s.net, "bill_count": s.bill_count,
        "revenue_refs": s.revenue_refs, "expense_refs": s.expense_refs,
        "price_changes": [{"product": c.product_name, "old": c.old_price, "new": c.new_price, "pct": c.pct,
                           "old_ref": c.old_ref, "new_ref": c.new_ref} for c in s.price_changes],
    }
    db.insert("daily_reports", {"report_date": day, "generated_at": now.isoformat(timespec="seconds"),
                                "text": text, "xlsx_path": str(path), "data_json": json.dumps(data, ensure_ascii=False)})
    db.commit()
    return text, path, s
