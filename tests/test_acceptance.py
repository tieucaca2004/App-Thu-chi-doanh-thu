"""Acceptance tests from the spec (§30) — checked on stored data and computed numbers, not just 'API = 200'."""
from openpyxl import load_workbook

from founder_assistant.analytics import daily_summary

from founder_assistant.report import generate_daily_report
from tests.conftest import extraction, item, revenue

from datetime import datetime
from zoneinfo import ZoneInfo

TZ = ZoneInfo("Asia/Ho_Chi_Minh")

BILL_OCR = """Thịt heo       5kg      450.000
Tôm            3kg      360.000
Rau                     120.000
Tổng                    930.000"""


def bill_extraction(**kw):
    return extraction(
        "PURCHASE_BILL", ocr_text=BILL_OCR, stated_total=930000, stated_total_text="930.000",
        items=[
            item("Thịt heo", 5, "kg", 450000, amount_text="450.000", evidence="Thịt heo       5kg      450.000"),
            item("Tôm", 3, "kg", 360000, amount_text="360.000", evidence="Tôm            3kg      360.000"),
            item("Rau", None, None, 120000, amount_text="120.000", evidence="Rau                     120.000"),
        ], **kw)


# ---------------------------------------------------------------- TEST 1 — VOICE
def test_1_voice_purchase(h):
    said = "Hôm nay mua 5 ký thịt heo 450 ngàn."
    res = h.send(extraction("PURCHASE_BILL", items=[
        item("thịt heo", 5, "ký", 450000, amount_text="450 ngàn", evidence="5 ký thịt heo 450 ngàn")]),
        audio=b"fake-aac-bytes", transcript=said)

    assert res.status == "confirmed"
    it = h.db.one("SELECT pi.*, p.name FROM purchase_items pi JOIN products p ON p.id = pi.product_id")
    assert it["name"] == "Thịt heo"
    assert it["quantity"] == 5 and it["unit"] == "kg" and it["amount"] == 450000
    assert it["amount_source"] == "source"
    assert it["unit_price"] == 90000 and it["unit_price_source"] == "computed"  # computed by code, not AI
    # raw data kept: voice file + transcript
    msg = h.db.one("SELECT * FROM messages")
    assert msg["transcript"] == said
    media = h.db.one("SELECT * FROM media")
    assert media["kind"] == "audio" and media["sha256"]
    assert "✅" in res.reply and "450.000đ" in res.reply


# ---------------------------------------------------------------- TEST 2 — BILL
def test_2_bill_image(h):
    res = h.send(bill_extraction(supplier="Chợ Bến Thành"), image=b"jpeg-bill-1")
    assert res.status == "confirmed"
    p = h.db.one("SELECT * FROM purchases")
    assert p["stated_total"] == 930000 and p["computed_total"] == 930000
    rows = h.db.all("SELECT raw_name, quantity, unit, amount, unit_price FROM purchase_items ORDER BY id")
    assert [(r["raw_name"], r["quantity"], r["unit"], r["amount"]) for r in rows] == [
        ("Thịt heo", 5, "kg", 450000), ("Tôm", 3, "kg", 360000), ("Rau", None, None, 120000)]
    assert rows[1]["unit_price"] == 120000
    assert "Tổng: 930.000đ" in res.reply and "3 mặt hàng" in res.reply
    # image kept, extraction kept, all linked
    ex = h.db.one("SELECT * FROM ai_extractions")
    assert ex["ocr_text"] == BILL_OCR and p["extraction_id"] == ex["id"]
    assert h.db.one("SELECT file_path FROM media")["file_path"].endswith(".jpg")


def test_2b_total_mismatch_is_held_back(h):
    ex = bill_extraction()
    ex.stated_total, ex.stated_total_text = 950000, "950.000"
    ex.ocr_text = BILL_OCR.replace("930.000", "950.000")
    res = h.send(ex, image=b"jpeg-bad-total")
    assert res.status == "needs_confirmation"
    assert "không khớp" in res.reply and "chưa đưa vào báo cáo" in res.reply
    text, _, s = generate_daily_report(h.db, "2026-09-26", h.settings.reports_dir, datetime(2026, 9, 26, 21, tzinfo=TZ))
    assert s.expense is None  # not counted until confirmed
    assert "M1 chờ xác nhận" in text


def test_2c_line_arithmetic_mismatch(h):
    ocr = "Thịt heo 5kg x 90.000 = 700.000"
    res = h.send(extraction(ocr_text=ocr, items=[
        item("Thịt heo", 5, "kg", 700000, unit_price=90000, amount_text="700.000", unit_price_text="90.000", evidence=ocr)]),
        image=b"jpeg-line-bad")
    assert res.status == "needs_confirmation"
    assert "Bill ghi Thịt heo 5kg × 90.000đ = 700.000đ. Tổng tiền không khớp phép tính" in res.reply


# ---------------------------------------------------------------- TEST 3 — PRICE CHANGE
def test_3_price_change(h):
    h.send(extraction(items=[item("thịt heo", 5, "kg", 440000, amount_text="440 ngàn", evidence="5 ký thịt heo 440 ngàn")]),
           text="mua 5 ký thịt heo 440 ngàn", when="2026-09-23 08:00")
    res = h.send(extraction(items=[item("thịt lợn", 5, "kg", 460000, amount_text="460 ngàn", evidence="5kg thịt lợn 460 ngàn")]),
                 text="mua 5kg thịt lợn 460 ngàn", when="2026-09-26 08:00")
    assert len(res.price_changes) == 1
    ch = res.price_changes[0]
    assert ch.product_name == "Thịt heo"  # 'thịt lợn' resolved to the same product
    assert (ch.old_price, ch.new_price, ch.diff) == (88000, 92000, 4000)
    assert round(ch.pct, 2) == 4.55
    assert "⚠️ Thịt heo tăng 4,55% so với lần mua trước" in res.reply
    alert = h.db.one("SELECT * FROM alerts WHERE code = 'PRICE_UP'")
    assert "88.000đ → 92.000đ/kg" in alert["message"]


def test_3b_units_never_compared_across_kinds(h):
    h.send(extraction(items=[item("Nước mắm", 1, "thùng", 300000, evidence="1 thùng nước mắm 300k", amount_text="300k")]),
           text="1 thùng nước mắm 300k")
    res = h.send(extraction(items=[item("Nước mắm", 2, "chai", 90000, evidence="2 chai nước mắm 90k", amount_text="90k")]),
                 text="2 chai nước mắm 90k")
    assert res.price_changes == []  # thùng vs chai: no comparison


def test_3c_gram_converted_to_kg(h):
    h.send(extraction(items=[item("Tỏi", 1, "kg", 60000, evidence="1kg tỏi 60k", amount_text="60k")]), text="1kg tỏi 60k")
    res = h.send(extraction(items=[item("Tỏi", 500, "g", 35000, evidence="500g tỏi 35k", amount_text="35k")]),
                 text="500g tỏi 35k")
    ch = res.price_changes[0]
    assert (ch.old_price, ch.new_price, ch.base_unit) == (60000, 70000, "kg")


def test_3d_unknown_unit_needs_confirmation_and_no_price_compare(h):
    res = h.send(extraction(items=[item("Tôm", 3, "rổ", 360000, evidence="3 rổ tôm 360k", amount_text="360k")]),
                 text="3 rổ tôm 360k")
    row = h.db.one("SELECT * FROM purchase_items")
    assert row["unit"] is None and row["status"] == "needs_confirmation"
    assert h.db.one("SELECT COUNT(*) c FROM price_history")["c"] == 0
    assert "Chưa rõ đơn vị: Tôm" in res.reply


# ---------------------------------------------------------------- TEST 4 — SALES
SALES_OCR = """Hủ tiếu xào hải sản x12  780.000
Hủ tiếu xào bò x8  480.000
Hủ tiếu xào thập cẩm x5  275.000
Tổng cộng 1.535.000"""


def sales_extraction():
    return extraction("SALES_BILL", ocr_text=SALES_OCR, stated_total=1535000, stated_total_text="1.535.000",
                      revenue=revenue(gross_revenue=1535000), items=[
                          item("Hủ tiếu xào hải sản", 12, None, 780000, amount_text="780.000", evidence="Hủ tiếu xào hải sản x12  780.000"),
                          item("Hủ tiếu xào bò", 8, None, 480000, amount_text="480.000", evidence="Hủ tiếu xào bò x8  480.000"),
                          item("Hủ tiếu xào thập cẩm", 5, None, 275000, amount_text="275.000", evidence="Hủ tiếu xào thập cẩm x5  275.000"),
                      ])


def test_4_sales_bill(h):
    res = h.send(sales_extraction(), image=b"jpeg-sales")
    assert res.status == "confirmed" and "Doanh thu: 1.535.000đ" in res.reply
    s = h.db.one("SELECT * FROM sales")
    assert s["revenue"] == 1535000 and s["revenue_source"] == "source"
    items = h.db.all("SELECT raw_name, quantity, unit_price FROM sale_items ORDER BY id")
    assert [(i["raw_name"], i["quantity"], i["unit_price"]) for i in items] == [
        ("Hủ tiếu xào hải sản", 12, 65000), ("Hủ tiếu xào bò", 8, 60000), ("Hủ tiếu xào thập cẩm", 5, 55000)]
    from founder_assistant.analytics import top_items_in_range
    top = top_items_in_range(h.db, "2026-09-26", "2026-09-26")
    assert [(n, q) for n, q, _ in top] == [("Hủ tiếu xào hải sản", 12), ("Hủ tiếu xào bò", 8), ("Hủ tiếu xào thập cẩm", 5)]


def test_4b_revenue_report_fields_not_invented(h):
    ocr = "CHỐT CA 26/09\nDoanh thu: 3.850.000\nSố bill: 42\nTiền mặt: 2.100.000\nChuyển khoản: 1.750.000"
    res = h.send(extraction("REVENUE_REPORT", ocr_text=ocr, document_date="2026-09-26", revenue=revenue(
        gross_revenue=3850000, bill_count=42, cash=2100000, bank_transfer=1750000)), image=b"jpeg-pos")
    assert res.status == "confirmed"
    s = h.db.one("SELECT * FROM sales")
    assert s["e_wallet"] is None and s["platform_fee"] is None  # absent in source -> NULL, not 0


# ---------------------------------------------------------------- TEST 5 — DAILY REPORT
def full_day(h):
    h.send(extraction(items=[item("Thịt heo", 5, "kg", 440000, evidence="5kg thịt heo 440k", amount_text="440k"),
                             item("Tôm", 3, "kg", 354000, evidence="3kg tôm 354k", amount_text="354k"),
                             item("Rau cải", 2, "kg", 56000, evidence="2kg rau cải 56k", amount_text="56k")]),
           text="5kg thịt heo 440k, 3kg tôm 354k, 2kg rau cải 56k", when="2026-09-23 08:00")
    ocr = ("Thịt heo 5kg 460.000\nTôm 3kg 375.000\nRau cải 2kg 48.000\nHành 47.000\nTổng 930.000")
    h.send(extraction(ocr_text=ocr, stated_total=930000, stated_total_text="930.000", supplier=None, items=[
        item("Thịt heo", 5, "kg", 460000, amount_text="460.000", evidence="Thịt heo 5kg 460.000"),
        item("Tôm", 3, "kg", 375000, amount_text="375.000", evidence="Tôm 3kg 375.000"),
        item("Rau cải", 2, "kg", 48000, amount_text="48.000", evidence="Rau cải 2kg 48.000"),
        item("Hành", None, None, 47000, amount_text="47.000", evidence="Hành 47.000")]),
        image=b"jpeg-market", when="2026-09-26 07:30")
    h.send(extraction("EXPENSE", items=[item("tiền gas", 1, "bình", 180000, category="gas",
                                             evidence="tiền gas 180 ngàn", amount_text="180 ngàn")]),
           text="trả tiền gas 180 ngàn", when="2026-09-26 10:00")
    ocr = "CHỐT CA\nDoanh thu 3.850.000\nSố bill 42"
    h.send(extraction("REVENUE_REPORT", ocr_text=ocr, revenue=revenue(gross_revenue=3850000, bill_count=42)),
           image=b"jpeg-pos", when="2026-09-26 20:00")
    h.send(sales_extraction(), image=b"jpeg-sales", when="2026-09-26 20:05")


def test_5_daily_report(h, monkeypatch):
    monkeypatch.setenv("PRIMARY_REVENUE_SOURCE", "pos_closing")  # Founder-confirmed authoritative source
    full_day(h)
    text, path, s = generate_daily_report(h.db, "2026-09-26", h.settings.reports_dir,
                                          datetime(2026, 9, 26, 21, 30, tzinfo=TZ))
    assert s.revenue == 3850000          # chốt ca only; sales bill not double counted
    assert s.expense == 1110000
    assert s.expense_by_category == {"ingredient": 930000, "gas": 180000}
    assert s.net == 2740000
    assert s.bill_count == 42
    for fragment in ["FOUNDER DAILY REPORT", "26/09/2026", "3.850.000đ", "1.110.000đ", "+2.740.000đ",
                     "chưa phải lợi nhuận ròng", "⚠️ Thịt heo", "88k → 92k/kg", "+4,55%",
                     "⚠️ Tôm", "118k → 125k/kg", "+5,93%", "↓ Rau cải", "28k → 24k/kg", "-14,29%",
                     "Hủ tiếu xào hải sản: 12 phần", "1 bill chưa xác định được nhà cung cấp",
                     "Chưa có dữ liệu chi phí điện tháng này", "chỉ dùng để đối chiếu, không cộng thêm",
                     "POS: 3.850.000đ; Bill evidence: 1.535.000đ; Chênh lệch: 2.315.000đ"]:
        assert fragment in text, fragment
    assert "lợi nhuận ròng:" not in text.lower()

    assert path.name == "Founder-Daily-Report-2026-09-26.xlsx"
    wb = load_workbook(path)
    assert wb.sheetnames == ["SUMMARY", "PURCHASE", "SALES", "PRICE_CHANGE", "ALERTS", "SOURCES"]
    summary = {r[0]: r[1] for r in wb["SUMMARY"].iter_rows(min_row=2, values_only=True)}
    assert summary["Doanh thu"] == 3850000 and summary["Tổng chi"] == 1110000
    assert summary["Thu - chi (chưa phải lợi nhuận ròng)"] == 2740000
    pc = {r[0]: r for r in wb["PRICE_CHANGE"].iter_rows(min_row=2, values_only=True)}
    assert pc["Tôm"][2:7] == (118000, 125000, 7000, 5.93, "Tăng")
    # traceability: "Tôm +5,93%" -> which bill, which message, which original file
    assert pc["Tôm"][9].startswith("M2 / msg#2") and pc["Tôm"][8].startswith("M1 / msg#1")
    sources = {r[0]: r for r in wb["SOURCES"].iter_rows(min_row=2, values_only=True)}
    # BUG-011: the POS closing (msg#4, source of the whole revenue) and the old price bill (msg#1) are listed
    assert {"msg#1", "msg#2", "msg#3", "msg#4", "msg#5"} <= set(sources)
    assert sources["msg#2"][4].endswith(".jpg") and len(sources["msg#2"][5]) == 64
    purchase = list(wb["PURCHASE"].iter_rows(min_row=2, values_only=True))
    hanh = next(r for r in purchase if r[2] == "Hành")
    assert hanh[4] == "UNKNOWN" and hanh[5] == "UNKNOWN"  # no quantity/unit in source -> UNKNOWN, not 0


def test_5b_empty_day_reports_unknown_not_zero(h):
    text, _, s = generate_daily_report(h.db, "2026-09-26", h.settings.reports_dir, datetime(2026, 9, 26, 21, tzinfo=TZ))
    assert s.revenue is None and s.expense is None and s.net is None
    assert "UNKNOWN — chưa có dữ liệu" in text


# ---------------------------------------------------------------- §16-17 revenue reconciliation
def _pos(h, amount_text, amount, when="2026-09-26 21:00"):
    ocr = f"CHỐT CA\nDoanh thu {amount_text}"
    return h.send(extraction("REVENUE_REPORT", ocr_text=ocr, revenue=revenue(gross_revenue=amount)),
                  image=f"pos-{amount}".encode(), when=when)


def _bill(h, name, qty, amount_text, amount, tag):
    line = f"{name} x{qty} {amount_text}"
    return h.send(extraction("SALES_BILL", ocr_text=line + f"\nTổng {amount_text}", stated_total=amount,
                             stated_total_text=amount_text, revenue=revenue(gross_revenue=amount),
                             items=[item(name, qty, None, amount, amount_text=amount_text, evidence=line)]),
                  image=tag.encode())


def test_revenue_pos_primary_equal_bills_not_doubled(h, monkeypatch):
    monkeypatch.setenv("PRIMARY_REVENUE_SOURCE", "pos_closing")
    _bill(h, "Hủ tiếu bò", 50, "3.000.000", 3000000, "b1")
    _bill(h, "Hủ tiếu hải sản", 30, "2.000.000", 2000000, "b2")
    _pos(h, "5.000.000", 5000000)
    s = daily_summary(h.db, "2026-09-26")
    assert s.revenue == 5000000  # not 10.000.000
    assert s.extra["reconciliation"]["matched"] is True
    assert not [a for a in s.alerts if a["code"].startswith("REVENUE_")]
    assert [n for n, _, _ in s.top_items] == ["Hủ tiếu bò", "Hủ tiếu hải sản"]  # bills still feed item stats


def test_revenue_mismatch_alert_numbers_untouched(h, monkeypatch):
    monkeypatch.setenv("PRIMARY_REVENUE_SOURCE", "pos_closing")
    _bill(h, "Hủ tiếu bò", 80, "4.850.000", 4850000, "b1")
    _pos(h, "5.000.000", 5000000)
    s = daily_summary(h.db, "2026-09-26")
    assert s.revenue == 5000000
    msg = next(a["message"] for a in s.alerts if a["code"] == "REVENUE_RECONCILIATION_MISMATCH")
    assert "POS: 5.000.000đ; Bill evidence: 4.850.000đ; Chênh lệch: 150.000đ. Cần Founder xác nhận" in msg


def test_revenue_bills_primary(h, monkeypatch):
    monkeypatch.setenv("PRIMARY_REVENUE_SOURCE", "sales_bills")
    _bill(h, "Hủ tiếu bò", 80, "4.850.000", 4850000, "b1")
    _pos(h, "5.000.000", 5000000)
    assert daily_summary(h.db, "2026-09-26").revenue == 4850000


def test_revenue_source_unset_is_unknown_not_guessed(h, monkeypatch):
    monkeypatch.delenv("PRIMARY_REVENUE_SOURCE", raising=False)
    _bill(h, "Hủ tiếu bò", 80, "4.850.000", 4850000, "b1")
    _pos(h, "5.000.000", 5000000)
    text, _, s = generate_daily_report(h.db, "2026-09-26", h.settings.reports_dir, datetime(2026, 9, 26, 22, tzinfo=TZ))
    assert s.revenue is None and s.net is None
    assert any(a["code"] == "REVENUE_SOURCE_UNSET" and a["severity"] == "critical" for a in s.alerts)
    assert "chưa xác định nguồn doanh thu chính" in text


def test_single_source_day_needs_no_config(h, monkeypatch):
    monkeypatch.delenv("PRIMARY_REVENUE_SOURCE", raising=False)
    _pos(h, "5.000.000", 5000000)
    s = daily_summary(h.db, "2026-09-26")
    assert s.revenue == 5000000 and s.extra["reconciliation"] is None
